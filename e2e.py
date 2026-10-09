"""Prueba de punta a punta del sistema Recorridas.

Levanta el servidor con una base vacía, y con un navegador real (Playwright/Chromium)
recorre todas las funciones del panel web y de la app del celular.
Uso:  python3 tests/e2e.py
"""
import io, json, os, re, shutil, socket, subprocess, sys, tempfile, threading, time, urllib.request
from http.server import BaseHTTPRequestHandler, HTTPServer

import openpyxl
from PIL import Image
from playwright.sync_api import sync_playwright, expect

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
WORK = tempfile.mkdtemp(prefix="recorridas_test_")
DATA = os.path.join(WORK, "data")
PORT = 4100
BASE = f"http://127.0.0.1:{PORT}"
RESULTS = []


def check(name, cond, detail=""):
    RESULTS.append((name, bool(cond), detail))
    print(("  OK   " if cond else "  FALLA ") + name + (f"  [{detail}]" if detail and not cond else ""))
    return cond


# ---------- servidor falso de links cortos de Maps (simula maps.app.goo.gl) ----------
class Redir(BaseHTTPRequestHandler):
    def do_GET(self):
        self.send_response(302)
        self.send_header("Location", "https://www.google.com/maps/place/Nuevo/@-34.8105,-55.9205,17z/data=!3d-34.81051!4d-55.92052")
        self.end_headers()

    def log_message(self, *a):
        pass


def free_port():
    s = socket.socket(); s.bind(("127.0.0.1", 0)); p = s.getsockname()[1]; s.close(); return p


# ---------- datos de prueba ----------
PINAR = [(-34.7950 + dx, -55.8900 + dy) for dx, dy in [(0.000, 0.000), (0.002, 0.001), (0.004, -0.002), (-0.003, 0.004), (0.006, 0.006),
                                                   (-0.005, -0.004), (0.001, 0.009), (0.008, -0.006), (-0.007, 0.007), (0.003, 0.012)]]
SOLY = [(-34.8230 + dx, -55.9550 + dy) for dx, dy in [(0, 0), (0.002, 0.003), (-0.003, 0.002), (0.004, -0.004)]]


def link(lat, lng):
    return f"https://www.google.com/maps?q={lat:.5f},{lng:.5f}"


def make_xlsx(path, rows, header=("Empresa", "Nombre", "Dirección", "Zona", "Horario", "Link Google Maps"), extra_template_rows=False):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Clientes"
    ws.append(list(header))
    if extra_template_rows:
        ws.append(["Obligatorio", "Obligatorio", "Obligatorio", "Obligatorio", "Opcional", "Obligatorio"])
        ws.append(["BEC", "Restaurante Ejemplo", "Av. Ejemplo 1234", "El Pinar", "", "https://www.google.com/maps?q=-34.7950,-55.8900"])
    for r in rows:
        ws.append(list(r))
    wb.save(path)


def jpg(path, color):
    Image.new("RGB", (2400, 1800), color).save(path, "JPEG")


def main():
    os.makedirs(DATA, exist_ok=True)
    rport = free_port()
    rsrv = HTTPServer(("127.0.0.1", rport), Redir)
    threading.Thread(target=rsrv.serve_forever, daemon=True).start()

    env = dict(os.environ, DATA_DIR=DATA, PORT=str(PORT), ADMIN_PASS="123456", VENDOR_PASS="1234")
    srv = subprocess.Popen(["node", "server.js"], cwd=ROOT, env=env, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
    for _ in range(50):
        try:
            urllib.request.urlopen(BASE + "/ingresar"); break
        except Exception:
            time.sleep(0.2)

    # Archivo 1: 10 en El Pinar, 4 en Solymar, filas de la plantilla, un repetido interno, dos con error
    f1 = os.path.join(WORK, "clientes1.xlsx")
    rows1 = []
    for i, (la, ln) in enumerate(PINAR):
        emp = "MAN" if i % 3 == 0 else "BEC"
        rows1.append((emp, f"{'Super' if emp == 'MAN' else 'Resto'} Pinar {i+1}", f"Calle {i+1}", "El Pinar", "10:00-12:00" if i == 4 else "", link(la, ln)))
    for i, (la, ln) in enumerate(SOLY):
        rows1.append(("bec", f"Resto Solymar {i+1}", f"Av Sol {i+1}", "Solymar", "", link(la, ln)))
    rows1.append(("BEC", "Resto Pinar 2", "Calle 2", "El Pinar", "", link(*PINAR[1])))  # repetida en el archivo
    rows1.append(("XX", "Empresa mala", "C", "El Pinar", "", link(-34.7, -55.8)))      # error empresa
    rows1.append(("BEC", "Sin link", "C", "El Pinar", "", ""))                         # error link
    make_xlsx(f1, rows1, extra_template_rows=True)

    # Archivo 2: 2 iguales, 1 con horario cambiado, 2 nuevas, 1 posible repetido
    f2 = os.path.join(WORK, "clientes2.csv")
    with open(f2, "w", encoding="utf-8") as fh:
        fh.write("Empresa;Nombre;Dirección;Zona;Horario;Link Google Maps\n")
        fh.write(f"MAN;Super Pinar 1;Calle 1;El Pinar;;{link(*PINAR[0])}\n")
        fh.write(f"BEC;Resto Pinar 3;Calle 3;El Pinar;;{link(*PINAR[2])}\n")
        fh.write(f"BEC;Resto Pinar 2;Calle 2;El Pinar;09:00-11:00;{link(*PINAR[1])}\n")
        fh.write(f"BEC;Resto Lagomar 1;Rambla 1;Lagomar;;{link(-34.8330, -55.9850)}\n")
        fh.write(f"MAN;Super Lagomar 1;Rambla 2;Lagomar;;{link(-34.8340, -55.9870)}\n")
        fh.write(f"BEC;Resto Pinar 3;Otra calle 99;El Pinar;;{link(-34.7700, -55.8600)}\n")

    img1, img2 = os.path.join(WORK, "f1.jpg"), os.path.join(WORK, "f2.jpg")
    jpg(img1, (120, 160, 120)); jpg(img2, (160, 120, 120))

    errors = []
    with sync_playwright() as pw:
        browser = pw.chromium.launch()

        # ================= ADMIN: panel web =================
        print("\nPanel web · administrador")
        ctx = browser.new_context(viewport={"width": 1366, "height": 900}, accept_downloads=True)
        page = ctx.new_page()
        page.on("pageerror", lambda e: errors.append(f"admin: {e}"))
        page.goto(BASE + "/panel")
        check("Sin sesión, /panel lleva a ingresar", "/ingresar" in page.url)
        page.fill("#u", "plateroadmin"); page.fill("#p", "mal"); page.click("#b")
        expect(page.locator("#e")).to_contain_text("incorrectos")
        check("Contraseña incorrecta muestra error", True)
        page.fill("#p", "123456"); page.click("#b")
        page.wait_for_url("**/panel**")
        expect(page.get_by_role("tab", name="Control de jornada")).to_be_visible()
        check("Administrador ve las 4 pestañas", page.get_by_role("tab").count() == 4)

        # Importación 1
        page.set_input_files("#xls", f1)
        expect(page.locator(".modal h2")).to_contain_text("Importar")
        txt = page.locator(".modal").inner_text()
        check("Vista previa: 14 nuevas", re.search(r"\b14\s*\n\s*Nuevas", txt) is not None, txt[:300])
        check("Vista previa: 2 filas con error", "Fila" in txt and "Empresa debe ser BEC o MAN" in txt and "Falta el link" in txt)
        check("Vista previa: detecta repetida dentro del archivo", "repetida dentro del archivo" in txt)
        check("Filas de ejemplo de la plantilla se ignoran", "Restaurante Ejemplo" not in txt)
        page.click("#imp-ok")
        expect(page.locator(".toast")).to_contain_text("14 nuevas")
        page.wait_for_selector("tbody tr")
        check("Tabla con 14 sucursales", page.locator("tbody tr[data-id]").count() == 14)

        # Importación 2 (CSV con ;)
        page.set_input_files("#xls", f2)
        expect(page.locator(".modal h2")).to_contain_text("Importar")
        txt = page.locator(".modal").inner_text()
        check("2º archivo: 2 nuevas", re.search(r"\b2\s*\n\s*Nuevas", txt) is not None, txt[:300])
        check("2º archivo: 1 se actualiza (horario)", re.search(r"\b1\s*\n\s*Se actualizan", txt) is not None and "horario" in txt)
        check("2º archivo: 2 ya estaban igual", re.search(r"\b2\s*\n\s*Ya estaban igual", txt) is not None)
        check("2º archivo: avisa posible repetido", "Posibles repetidos" in txt)
        page.click("#imp-ok")
        expect(page.locator(".toast")).to_contain_text("2 nuevas, 1 actualizadas")
        page.wait_for_timeout(300)
        n_rows = page.locator("tbody tr[data-id]").count()
        check("Se acumulan sin duplicar: 16 sucursales", n_rows == 16, str(n_rows))
        hor = page.locator("tr:has(input[value='Resto Pinar 2']) input[data-f='horario']").input_value()
        check("Horario actualizado por el Excel", hor == "09:00–11:00", hor)

        # Alta manual con link corto (redirección) y duplicado
        page.click("text=+ Agregar sucursal")
        page.select_option(".modal select[name=empresa]", "MAN")
        page.fill(".modal input[name=nombre]", "Super Link Corto")
        page.fill(".modal input[name=direccion]", "Ruta 1")
        page.fill(".modal input[name=zona]", "El Pinar")
        page.fill(".modal input[name=link]", f"http://127.0.0.1:{rport}/abc")
        page.click("#fn-b")
        expect(page.locator(".toast")).to_contain_text("agregada")
        page.wait_for_timeout(300)
        check("Alta manual con link corto de Maps", page.locator("input[value='Super Link Corto']").count() == 1)
        page.click("text=+ Agregar sucursal")
        page.fill(".modal input[name=nombre]", "Otro nombre")
        page.fill(".modal input[name=zona]", "El Pinar")
        page.fill(".modal input[name=link]", link(*PINAR[0]))
        page.click("#fn-b")
        expect(page.locator("#fn-e")).to_contain_text("ya está cargada")
        check("Alta manual duplicada se rechaza", True)
        page.click(".modal [data-close]")

        # Edición en la tabla
        inp = page.locator("tr:has(input[value='Resto Pinar 3']) input[data-f='direccion']")
        inp.fill("Calle 3 esquina 4"); inp.press("Tab")
        page.wait_for_timeout(500)
        cl = json.loads(page.evaluate("fetch('/api/clients').then(r=>r.text())"))
        check("Editar celda guarda en la base", any(c["direccion"] == "Calle 3 esquina 4" for c in cl))
        lk = page.locator("tr:has(input[value='Resto Pinar 3']) input[data-f='link']")
        old = lk.input_value()
        lk.fill("esto no es un link"); lk.press("Tab")
        expect(page.locator(".toast")).to_contain_text("No se pudo leer")
        check("Link inválido se rechaza y vuelve al valor anterior", lk.input_value() == old)

        # Eliminar
        page.locator("tr:has(input[value='Super Lagomar 1']) [data-act=borrar]").click()
        page.click(".modal [data-r='1']")
        expect(page.locator(".toast")).to_contain_text("eliminada")
        check("Eliminar sucursal", page.locator("tbody tr[data-id]").count() == 16)

        # Filtros
        page.click("button.chipb[data-z='Solymar']")
        check("Filtro por zona", page.locator("tbody tr[data-id]").count() == 4)
        page.click("button.chipb[data-z='Todas']")
        page.fill("#bq", "lagomar")
        page.wait_for_timeout(200)
        check("Búsqueda", page.locator("tbody tr[data-id]").count() == 1)

        # Plantilla
        with page.expect_download() as d:
            page.click("text=Descargar plantilla")
        p_tpl = os.path.join(WORK, "plantilla.xlsx"); d.value.save_as(p_tpl)
        wb = openpyxl.load_workbook(p_tpl)
        check("Plantilla descargable con encabezados", [c.value for c in wb["Clientes"][1]][:6] == ["Empresa", "Nombre", "Dirección", "Zona", "Horario", "Link Google Maps"])
        ctx.close()

        # ================= VENDEDOR: app del celular =================
        print("\nApp del celular · vendedor")
        iphone = pw.devices["iPhone 13"]
        mctx = browser.new_context(**iphone, geolocation={"latitude": -34.7990, "longitude": -55.8990}, permissions=["geolocation"])
        m = mctx.new_page()
        m.on("pageerror", lambda e: errors.append(f"app: {e}"))
        m.goto(BASE + "/")
        m.fill("#u", "msayago"); m.fill("#p", "1234"); m.click("#b")
        m.wait_for_url("**/app")
        expect(m.locator("h1")).to_have_text("¿A qué zona vas hoy?")
        check("Vendedor entra a la app del celular", True)
        zones_txt = m.locator("main").inner_text()
        check("Muestra zonas con avance del mes", "El Pinar" in zones_txt and "Solymar" in zones_txt and "0/11 este mes" in zones_txt, zones_txt[:200])
        m.click("button.zone[data-z='El Pinar']")
        for _ in range(10):
            m.click("[data-act=menos]")
        for _ in range(4):
            m.click("[data-act=mas]")
        expect(m.locator("#n")).to_have_text("5")
        expect(m.locator("main")).to_contain_text(re.compile(r"Tiempo estimado: \d"))
        est = m.locator("main").inner_text()
        check("Estimación de tiempo de la recorrida", re.search(r"Tiempo estimado: \d+ h \d\d|Tiempo estimado: \d+ min", est) is not None, est[-120:])
        m.click("[data-act=armar]")
        expect(m.locator("h1")).to_have_text("El Pinar")
        stops = m.locator("button.stop")
        check("Recorrida con 5 paradas", stops.count() == 5)
        cur = json.loads(m.evaluate("fetch('/api/plan/current').then(r=>r.text())"))
        # el primero tiene que ser el más cercano a la ubicación del vendedor
        import math
        def d(a, b): return math.hypot((a[0]-b[0])*111, (a[1]-b[1])*91.5)
        start = (-34.7990, -55.8990)
        first = (cur["items"][0]["client"]["lat"], cur["items"][0]["client"]["lng"])
        others = [(c["client"]["lat"], c["client"]["lng"]) for c in cur["items"][1:]]
        check("Primera parada = la más cercana al vendedor", all(d(start, first) <= d(start, o) + 1e-9 for o in others))
        check("Plan guarda ubicación de inicio", cur["plan"]["start_lat"] == -34.799)
        check("Link de ruta en Google Maps", m.locator("a:has-text('Ruta en Maps')").first.get_attribute("href").startswith("https://www.google.com/maps/dir/-34.799,-55.899/"))
        names = [it["client"]["nombre"] for it in cur["items"]]

        # Parada 1: SÍ + foto + comentario
        stops.nth(0).click()
        c0 = cur["items"][0]["client"]
        mctx.set_geolocation({"latitude": c0["lat"] + 0.0001, "longitude": c0["lng"]})
        guardar = m.locator("[data-act=guardar]")
        check("Guardar deshabilitado antes de marcar SÍ/NO", guardar.is_disabled())
        m.click("button.s")
        expect(m.locator(".reg")).to_contain_text("Registrado automáticamente")
        check("SÍ registra hora y ubicación", re.search(r"\d\d:\d\d · -34\.79", m.locator(".reg").inner_text()) is not None)
        check("SÍ sin foto no permite guardar", guardar.is_disabled())
        title = m.locator("section.sec h2").nth(1).inner_text()
        check("Pide foto de góndola (MAN) o del local (BEC)", title == ("Foto de la góndola" if c0["empresa"] == "MAN" else "Foto del local"), title)
        m.set_input_files("input[data-act=foto]", img1)
        expect(m.locator("img.photo")).to_be_visible()
        check("Foto subida y visible", True)
        m.click("button.chip:has-text('Cliente interesado')")
        m.fill("#nota", "Cliente interesado. Pidió lista de precios")
        expect(guardar).to_be_enabled()
        guardar.click()
        expect(m.locator(".toast")).to_contain_text("Visita guardada")
        check("Visita SÍ guardada", "Visita SÍ" in m.locator("button.stop").nth(0).inner_text())

        # Parada 2: NO (foto lejos no aplica), comentario
        m.locator("button.stop").nth(1).click()
        m.click("button.n")
        expect(m.locator("main")).to_contain_text("vuelve a las sugerencias")
        m.click("button.chip:has-text('No estaba el encargado')")
        m.click("[data-act=guardar]")
        expect(m.locator(".toast")).to_contain_text("Visita guardada")
        check("Visita NO guardada", "Visita NO" in m.locator("button.stop").nth(1).inner_text())

        # Parada 3: SÍ con foto sacada lejos (para la alerta)
        m.locator("button.stop").nth(2).click()
        c2 = cur["items"][2]["client"]
        mctx.set_geolocation({"latitude": c2["lat"], "longitude": c2["lng"] + 0.0002})
        m.click("button.s")
        expect(m.locator(".reg")).to_be_visible()
        mctx.set_geolocation({"latitude": c2["lat"] + 0.02, "longitude": c2["lng"]})
        m.set_input_files("input[data-act=foto]", img2)
        expect(m.locator("img.photo")).to_be_visible()
        m.click("[data-act=guardar]")
        expect(m.locator(".toast")).to_contain_text("Visita guardada")

        # Recargar la página no pierde la recorrida
        m.reload()
        expect(m.locator("h1")).to_have_text("El Pinar")
        check("Al recargar sigue la recorrida en curso", m.locator("button.stop").count() == 5)

        # Terminar jornada
        mctx.set_geolocation({"latitude": -34.8102, "longitude": -55.9312})
        m.click("[data-act=terminar]")
        m.click(".modal [data-r='1']")
        expect(m.locator(".err")).to_contain_text("Jornada terminada")
        fin = m.locator(".err").inner_text()
        check("Terminar jornada registra ubicación", "Ubicación de cierre registrada" in fin, fin)
        check("Las 2 no visitadas vuelven a sugerirse", "2 sucursales vuelven" in fin, fin)
        pinar_txt = m.locator("button.zone[data-z='El Pinar']").inner_text()
        check("Zona muestra 2 visitadas y 3 pendientes", "2/11 este mes" in pinar_txt and "3 quedaron pendientes" in pinar_txt, pinar_txt)

        # Nueva recorrida en la misma zona: primero las pendientes, nunca las visitadas
        m.click("button.zone[data-z='El Pinar']")
        m.click("[data-act=armar]")
        expect(m.locator("h1")).to_have_text("El Pinar")
        cur2 = json.loads(m.evaluate("fetch('/api/plan/current').then(r=>r.text())"))
        n2 = [it["client"]["nombre"] for it in cur2["items"]]
        visited = {names[0], names[2]}
        postponed = {names[1], names[3], names[4]}
        check("No repite sucursales visitadas en el mes", not (visited & set(n2)), str(n2))
        check("Incluye todas las que quedaron pendientes", postponed <= set(n2), str(n2))
        check("Marca 'Quedó pendiente la vez anterior'", m.locator("text=Quedó pendiente la vez anterior").count() == 3)

        # Armar otra sin terminar la jornada: la anterior se cierra sola
        m.click("[data-act=salir]"); m.click(".modal [data-r='0']")
        m.evaluate("fetch('/api/plan',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({zona:'Solymar',n:2,lat:-34.823,lng:-55.955})})")
        m.wait_for_timeout(400)
        m.reload()
        expect(m.locator("h1")).to_have_text("Solymar")
        check("Jornada sin cerrar se cierra sola al armar otra", True)
        # sesión segura: fotos protegidas
        r = urllib.request.Request(BASE + "/api/visits")
        try:
            urllib.request.urlopen(r); anon = 200
        except urllib.error.HTTPError as e:
            anon = e.code
        check("Sin sesión no se accede a los datos", anon == 401)
        mctx.close()

        # ================= VENDEDOR en la web =================
        print("\nPanel web · vendedor")
        vctx = browser.new_context(viewport={"width": 1280, "height": 900})
        v = vctx.new_page()
        v.on("pageerror", lambda e: errors.append(f"vend-web: {e}"))
        v.goto(BASE + "/ingresar")
        v.fill("#u", "msayago"); v.fill("#p", "1234"); v.click("#b")
        v.wait_for_url("**/panel**")
        expect(v.get_by_role("tab", name="Clientes")).to_be_visible()
        check("Vendedor en la web no ve Control de jornada", v.get_by_role("tab", name="Control de jornada").count() == 0)
        status = v.evaluate("fetch('/api/control').then(r=>r.status)")
        check("API de control bloqueada para el vendedor", status == 403)
        vis = json.loads(v.evaluate("fetch('/api/visits').then(r=>r.text())"))
        check("Vendedor no recibe datos de ubicación ni tiempos", vis and all("lat" not in x and "minutos" not in x for x in vis))
        v.click("[role=tab]:has-text('Visitas')")
        expect(v.locator("h1")).to_have_text("Visitas")
        check("Vendedor ve sus visitas (3)", v.locator("tbody tr").count() == 3)
        v.click("[role=tab]:has-text('Clientes')")
        expect(v.locator("text=+ Agregar sucursal")).to_be_visible()
        check("Vendedor puede cargar sucursales desde la web", v.locator("#xls").count() == 1)
        # cambio de contraseña
        v.click("[data-act=clave]")
        v.fill(".modal input[name=current]", "1234"); v.fill(".modal input[name=next]", "nueva123")
        v.click(".modal button.dark")
        expect(v.locator(".toast")).to_contain_text("Contraseña cambiada")
        v.click("[data-act=salir]")
        v.wait_for_url("**/ingresar")
        v.fill("#u", "msayago"); v.fill("#p", "1234"); v.click("#b")
        expect(v.locator("#e")).to_be_visible()
        v.fill("#p", "nueva123"); v.click("#b")
        v.wait_for_url("**/panel**")
        check("Cambio de contraseña funciona", True)
        v.evaluate("fetch('/api/password',{method:'POST',headers:{'Content-Type':'application/json'},body:JSON.stringify({current:'nueva123',next:'123456'})})")
        vctx.close()

        # ================= ADMIN: zonas, visitas, control =================
        print("\nPanel web · análisis y control (administrador)")
        ctx = browser.new_context(viewport={"width": 1366, "height": 900}, accept_downloads=True)
        a = ctx.new_page()
        a.on("pageerror", lambda e: errors.append(f"admin2: {e}"))
        a.goto(BASE + "/ingresar")
        a.fill("#u", "plateroadmin"); a.fill("#p", "123456"); a.click("#b")
        a.wait_for_url("**/panel**")
        a.click("[role=tab]:has-text('Zonas del mes')")
        a.wait_for_selector("#map .leaflet-interactive")
        check("Zonas: mapa con las sucursales", a.locator("#map path.leaflet-interactive").count() == 11)
        ztxt = a.locator("button.zcard[data-z='El Pinar']").inner_text()
        check("Zonas: avance de El Pinar 2/11", "2/11" in ztxt, ztxt)
        check("Zonas: tabla con estados", a.locator("text=Quedó pendiente · vuelve primero").count() >= 3)

        a.click("[role=tab]:has-text('Visitas')")
        expect(a.locator("h1")).to_have_text("Visitas")
        check("Visitas: 3 registros", a.locator("tbody tr").count() == 3)
        check("Visitas: admin ve ubicación y foto", a.locator("text=Ver en mapa").count() == 3 and a.locator("button.thumb").count() == 2)
        a.locator("button.thumb").first.click()
        expect(a.locator(".lb img")).to_be_visible()
        ok_img = a.evaluate("document.querySelector('.lb img').naturalWidth")
        check("Foto se abre en grande (comprimida a 1600 px)", ok_img == 1600, str(ok_img))
        a.click(".lb")
        a.select_option("#vr", "no")
        a.wait_for_timeout(400)
        check("Filtro de resultado", a.locator("tbody tr").count() == 1)
        a.select_option("#vr", "")
        a.wait_for_timeout(400)
        with a.expect_download() as d:
            a.click("text=Exportar a Excel")
        px = os.path.join(WORK, "visitas.xlsx"); d.value.save_as(px)
        ws = openpyxl.load_workbook(px).active
        hdr = [c.value for c in ws[1]]
        check("Exportar a Excel con todas las columnas", ws.max_row == 4 and "Minutos en el local" in hdr and "Distancia foto a la sucursal (m)" in hdr, str(hdr))

        a.click("[role=tab]:has-text('Control de jornada')")
        a.wait_for_selector(".ev")
        ctxt = a.locator(".wrap").inner_text()
        check("Control: inicio y fin de jornada", "Inicio de jornada · El Pinar" in ctxt and "Fin de jornada" in ctxt)
        check("Control: ubicación de fin de jornada", "-34.81020, -55.93120" in ctxt)
        check("Control: hora y distancia de cada foto", re.search(r"Foto \d\d:\d\d · -34\.", ctxt) is not None and "de la sucursal" in ctxt)
        check("Control: alerta de foto lejos de la sucursal", "la foto se sacó a 2,2 km" in ctxt, ctxt[ctxt.find('Alertas'):ctxt.find('Alertas')+400])
        check("Control: alerta de jornada que no se cerró", "No tocó “Terminar jornada”" in ctxt)
        check("Control: resumen de tiempos", "Jornada" in ctxt and "En visitas" in ctxt and "Traslados" in ctxt and "Sin actividad" in ctxt)
        check("Control: mapa del recorrido", a.locator("#map path.leaflet-interactive").count() >= 4)
        check("Control: tiempos estimados por visita", "Tiempo estimado por visita" in ctxt)
        a.fill("#cf", "2020-01-01"); a.dispatch_event("#cf", "change")
        expect(a.locator(".wrap")).to_contain_text("No hubo jornada ese día")
        check("Control: día sin jornada", True)
        ctx.close()
        browser.close()

    srv.terminate()
    out = srv.communicate(timeout=5)[0]
    check("Sin errores de JavaScript en las páginas", not errors, "; ".join(errors))
    check("Sin errores en el servidor", "Error" not in out, out[-500:])
    rsrv.shutdown()
    shutil.rmtree(WORK, ignore_errors=True)
    fails = [r for r in RESULTS if not r[1]]
    print(f"\n{len(RESULTS) - len(fails)}/{len(RESULTS)} pruebas OK")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
